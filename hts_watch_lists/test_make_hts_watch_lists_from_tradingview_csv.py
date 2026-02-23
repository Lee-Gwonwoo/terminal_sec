from __future__ import annotations

import argparse
import json
import os
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Final
from typing import Iterable

import pandas as pd
import requests
from openpyxl import load_workbook


@dataclass(frozen=True)
class Paths:
    template_xlsx: Path
    input_csv: Path
    out_dir: Path


def _retry(
    fn,
    *,
    tries: int = 10,
    initial_sleep_s: float = 0.25,
    backoff: float = 1.6,
    retry_on: tuple[type[BaseException], ...] = (Exception,),
    label: str,
):
    sleep_s = initial_sleep_s
    last_exc: BaseException | None = None
    for attempt in range(1, tries + 1):
        try:
            return fn()
        except retry_on as exc:  # noqa: PERF203
            last_exc = exc
            if attempt >= tries:
                break
            print(f"[{label}] retry {attempt}/{tries} failed: {exc!r} (sleep {sleep_s:.2f}s)")
            time.sleep(sleep_s)
            sleep_s *= backoff
    assert last_exc is not None
    raise last_exc


def _safe_replace(src_tmp: Path, dst: Path) -> None:
    def _do() -> None:
        os.replace(src_tmp, dst)

    _retry(_do, label=f"os.replace({dst.name})")


def _read_csv_symbols(csv_path: Path) -> list[str]:
    encodings_to_try = ["utf-8-sig", "utf-8", "cp949"]
    last_exc: Exception | None = None
    df = None
    for enc in encodings_to_try:
        try:
            df = pd.read_csv(csv_path, encoding=enc)
            last_exc = None
            break
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
    if df is None:
        assert last_exc is not None
        raise last_exc

    col_candidates = ["Symbol", "symbol", "ticker", "Ticker"]
    symbol_col = next((c for c in col_candidates if c in df.columns), None)
    if symbol_col is None:
        raise ValueError(f"CSV missing symbol column; looked for {col_candidates}. columns={list(df.columns)}")

    raw = df[symbol_col].astype(str).tolist()
    symbols: list[str] = []
    seen: set[str] = set()
    for s in raw:
        s = s.strip()
        if not s or s.lower() == "nan":
            continue

        # Allow TradingView style: NASDAQ:MSFT, NYSE:T
        if ":" in s and len(s.split(":", 1)[0]) <= 10:
            _, maybe_symbol = s.split(":", 1)
            if maybe_symbol:
                s = maybe_symbol.strip()

        if s not in seen:
            seen.add(s)
            symbols.append(s)

    return symbols


def _ensure_template_readable(template_path: Path) -> Path:
    """Return a readable XLSX path.

    The provided template in this repo may be truncated by 1 byte, which breaks
    ZIP parsing for openpyxl/zipfile. If that exact truncation is detected,
    create a _fixed copy by appending a single 0x00 byte.
    """
    import zipfile

    if zipfile.is_zipfile(template_path):
        return template_path

    b = template_path.read_bytes()
    sig = b"PK\x05\x06"
    idx = b.rfind(sig)
    if idx < 0:
        raise ValueError(f"Template is not a readable .xlsx zip: {template_path}")

    remaining = len(b) - idx
    if remaining == 21:
        fixed = template_path.with_name(template_path.stem + "_fixed" + template_path.suffix)
        if not fixed.exists():
            fixed.write_bytes(b + b"\x00")
            print(f"[template] wrote repaired copy: {fixed}")
        if not zipfile.is_zipfile(fixed):
            raise ValueError(f"Template repair attempt failed; still not a zip: {fixed}")
        return fixed

    raise ValueError(
        f"Template looks like zip but is not readable by zipfile (EOCD remaining bytes={remaining}); "
        f"please re-save it in Excel: {template_path}"
    )


def _find_header_cell(ws, header_text: str) -> tuple[int, int]:
    for r in range(1, 101):
        for c in range(1, 51):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() == header_text:
                return r, c
    raise ValueError(f"Could not find header '{header_text}' in worksheet '{ws.title}'")


def _load_exchange_cache(cache_path: Path) -> dict[str, str]:
    if not cache_path.exists():
        return {}
    try:
        return json.loads(cache_path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {}


def _save_exchange_cache(cache_path: Path, cache: dict[str, str]) -> None:
    tmp = cache_path.with_suffix(cache_path.suffix + ".tmp")
    tmp.write_text(json.dumps(cache, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    _safe_replace(tmp, cache_path)


def _batched(items: list[str], batch_size: int) -> Iterable[list[str]]:
    for i in range(0, len(items), batch_size):
        yield items[i : i + batch_size]

NASDAQ_TRADER_NASDAQ_LISTED_URL: Final[str] = "https://www.nasdaqtrader.com/dynamic/SymDir/nasdaqlisted.txt"
NASDAQ_TRADER_OTHER_LISTED_URL: Final[str] = "https://www.nasdaqtrader.com/dynamic/SymDir/otherlisted.txt"


def _download_text(url: str, *, timeout_s: int = 30) -> str:
    def _do() -> str:
        r = requests.get(url, timeout=timeout_s)
        r.raise_for_status()
        return r.text

    return _retry(_do, label=f"download:{url.split('/')[-1]}")


def _load_or_download_listing(out_dir: Path, *, filename: str, url: str, refresh: bool) -> str:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / filename
    if path.exists() and not refresh:
        return path.read_text(encoding="utf-8", errors="replace")

    print(f"[listing] downloading {filename} from nasdaqtrader.com")
    text = _download_text(url)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    _safe_replace(tmp, path)
    return text


def _parse_nasdaqlisted(text: str) -> set[str]:
    # Pipe-delimited. Footer: "File Creation Time:"
    lines = [ln.strip("\r") for ln in text.splitlines() if ln.strip()]
    out: set[str] = set()
    for i, ln in enumerate(lines):
        if i == 0:
            continue
        if ln.startswith("File Creation Time"):
            break
        sym = ln.split("|", 1)[0].strip()
        if sym and sym != "Symbol":
            out.add(sym)
    return out


def _parse_otherlisted(text: str) -> dict[str, str]:
    # Pipe-delimited. Header includes: ACT Symbol | ... | Exchange | ...
    lines = [ln.strip("\r") for ln in text.splitlines() if ln.strip()]
    if not lines:
        return {}

    header = lines[0].split("|")
    try:
        sym_idx = header.index("ACT Symbol")
        exch_idx = header.index("Exchange")
    except ValueError as exc:
        raise ValueError(f"Unexpected otherlisted header: {header[:10]}") from exc

    out: dict[str, str] = {}
    for ln in lines[1:]:
        if ln.startswith("File Creation Time"):
            break
        parts = ln.split("|")
        if len(parts) <= max(sym_idx, exch_idx):
            continue
        sym = parts[sym_idx].strip()
        exch = parts[exch_idx].strip()
        if sym and exch:
            out[sym] = exch
    return out


def _prefix_from_listings(symbol: str, *, nasdaq_symbols: set[str], other_exch: dict[str, str]) -> str:
    if symbol in nasdaq_symbols:
        return "ND"
    # otherlisted exchange codes: N=NYSE, A=NYSE American, P=NYSE Arca, Z=Bats, V=IEXG, ...
    # User request only distinguishes NASDAQ vs New York, so anything not NASDAQ is treated as NY.
    if symbol in other_exch:
        return "NY"
    return "NY"


def _resolve_prefixes(
    symbols: list[str],
    *,
    cache_path: Path,
    enable_lookup: bool,
    refresh_listings: bool = False,
) -> dict[str, str]:
    cache = _load_exchange_cache(cache_path)

    missing = [s for s in symbols if s not in cache]
    if missing and enable_lookup:
        out_dir = cache_path.parent
        nas_text = _load_or_download_listing(
            out_dir, filename="nasdaqlisted.txt", url=NASDAQ_TRADER_NASDAQ_LISTED_URL, refresh=refresh_listings
        )
        oth_text = _load_or_download_listing(
            out_dir, filename="otherlisted.txt", url=NASDAQ_TRADER_OTHER_LISTED_URL, refresh=refresh_listings
        )
        nasdaq_symbols = _parse_nasdaqlisted(nas_text)
        other_exch = _parse_otherlisted(oth_text)

        print(f"[exchange] resolving {len(missing)} missing symbols using NASDAQ Trader listings")
        for sym in missing:
            cache[sym] = _prefix_from_listings(sym, nasdaq_symbols=nasdaq_symbols, other_exch=other_exch)
        _save_exchange_cache(cache_path, cache)
    elif missing and not enable_lookup:
        print(f"[exchange] lookup disabled; defaulting {len(missing)} symbols to NY")
        for sym in missing:
            cache[sym] = "NY"
        _save_exchange_cache(cache_path, cache)

    return {sym: cache.get(sym, "NY") for sym in symbols}


def _write_chunk_xlsx(
    *,
    template_path: Path,
    out_path: Path,
    symbols: list[str],
    prefixes: dict[str, str],
) -> None:
    wb = load_workbook(template_path)
    ws = wb.active

    header_row, code_col = _find_header_cell(ws, "종목코드")
    start_row = header_row + 1

    # Clear existing values under 종목코드 (keep formatting)
    for r in range(start_row, start_row + max(len(symbols), 2000)):
        if ws.cell(r, code_col).value is None and r > start_row + len(symbols):
            break
        ws.cell(r, code_col).value = None

    for i, sym in enumerate(symbols):
        pref = prefixes.get(sym, "NY")
        ws.cell(start_row + i, code_col).value = f"{pref}{sym}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(out_path.suffix + ".tmp")
    wb.save(tmp)
    _safe_replace(tmp, out_path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate HTS watchlist Excel files from TradingView screener CSV, 100 symbols per file."
    )
    parser.add_argument(
        "--template",
        default=r"C:\github_coding\terminal_sec\hts_watch_lists\hts watch lists.xlsx",
        help="Template XLSX path (format source).",
    )
    parser.add_argument(
        "--csv",
        default=r"C:\github_coding\terminal_sec\tradigview_screener\watch lists2_2026-02-22.csv",
        help="Input TradingView CSV path.",
    )
    parser.add_argument(
        "--out-dir",
        default=r"C:\github_coding\terminal_sec\hts_watch_lists",
        help="Output directory where hts_watch_lists_*.xlsx will be written.",
    )
    parser.add_argument("--chunk-size", type=int, default=100)
    parser.add_argument("--start-index", type=int, default=1)
    parser.add_argument(
        "--no-exchange-lookup",
        action="store_true",
        help="Do not download listing files for exchange lookup; defaults all prefixes to NY.",
    )
    parser.add_argument(
        "--refresh-listings",
        action="store_true",
        help="Re-download nasdaqlisted.txt/otherlisted.txt even if cached in out-dir.",
    )
    parser.add_argument(
        "--cache",
        default="exchange_cache.json",
        help="Exchange cache JSON filename inside out-dir.",
    )

    args = parser.parse_args()
    paths = Paths(template_xlsx=Path(args.template), input_csv=Path(args.csv), out_dir=Path(args.out_dir))

    if args.chunk_size <= 0:
        raise ValueError("--chunk-size must be > 0")
    if args.start_index <= 0:
        raise ValueError("--start-index must be >= 1")

    template_path = _ensure_template_readable(paths.template_xlsx)
    symbols = _read_csv_symbols(paths.input_csv)
    print(f"[input] symbols: {len(symbols)}")

    cache_path = paths.out_dir / args.cache
    prefixes = _resolve_prefixes(
        symbols,
        cache_path=cache_path,
        enable_lookup=not args.no_exchange_lookup,
        refresh_listings=bool(args.refresh_listings),
    )

    chunks = list(_batched(symbols, args.chunk_size))
    print(f"[output] chunks: {len(chunks)} (chunk_size={args.chunk_size})")

    for i, chunk in enumerate(chunks, start=args.start_index):
        out_path = paths.out_dir / f"hts_watch_lists_{i}.xlsx"
        _write_chunk_xlsx(template_path=template_path, out_path=out_path, symbols=chunk, prefixes=prefixes)
        print(f"[output] wrote: {out_path} ({len(chunk)} symbols)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
