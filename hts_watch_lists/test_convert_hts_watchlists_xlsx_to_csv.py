from __future__ import annotations

import argparse
import os
import time
from pathlib import Path

from openpyxl import load_workbook


def _retry(
    fn,
    *,
    tries: int = 10,
    initial_sleep_s: float = 0.25,
    backoff: float = 1.6,
    label: str,
):
    sleep_s = initial_sleep_s
    last_exc: BaseException | None = None
    for attempt in range(1, tries + 1):
        try:
            return fn()
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
            if attempt >= tries:
                break
            print(f"[{label}] retry {attempt}/{tries} failed: {exc!r} (sleep {sleep_s:.2f}s)")
            time.sleep(sleep_s)
            sleep_s *= backoff
    assert last_exc is not None
    raise last_exc


def _safe_replace(src_tmp: Path, dst: Path) -> None:
    def _do() -> None:
        os.replace(src_tmp, dst)

    _retry(_do, label=f"os.replace({dst.name})")


def _find_header_cell(ws, header_text: str) -> tuple[int, int]:
    for r in range(1, 101):
        for c in range(1, 51):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() == header_text:
                return r, c
    raise ValueError(f"Could not find header '{header_text}' in worksheet '{ws.title}'")


def _extract_codes_from_xlsx(path: Path) -> list[str]:
    wb = load_workbook(path)
    ws = wb.active
    header_row, code_col = _find_header_cell(ws, "종목코드")
    start_row = header_row + 1

    codes: list[str] = []
    empty_after_data = 0
    for r in range(start_row, start_row + 20000):
        v = ws.cell(r, code_col).value
        if v is None or (isinstance(v, str) and not v.strip()):
            if codes:
                empty_after_data += 1
                if empty_after_data >= 3:
                    break
            continue
        empty_after_data = 0
        codes.append(str(v).strip())

    return codes


def _write_hts_csv(out_path: Path, codes: list[str], *, encoding: str) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(out_path.suffix + ".tmp")

    # Match the observed format in hts_watch_lists/1.csv
    lines: list[str] = ["Version=1.0,", "거래소,종목코드"]
    lines.extend([f",{code}" for code in codes])
    text = "\r\n".join(lines) + "\r\n"

    tmp.write_text(text, encoding=encoding, newline="")
    _safe_replace(tmp, out_path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert generated HTS watchlist XLSX files (hts_watch_lists_*.xlsx) to HTS CSV format (like hts_watch_lists/1.csv)."
    )
    parser.add_argument(
        "--dir",
        default=r"C:\github_coding\terminal_sec\hts_watch_lists",
        help="Directory containing hts_watch_lists_*.xlsx.",
    )
    parser.add_argument(
        "--pattern",
        default="hts_watch_lists_*.xlsx",
        help="Glob pattern for input XLSX files.",
    )
    parser.add_argument(
        "--encoding",
        default="cp949",
        help="Output encoding (1.csv in this repo appears to be cp949).",
    )
    args = parser.parse_args()

    out_dir = Path(args.dir)
    xlsx_paths = sorted(out_dir.glob(args.pattern), key=lambda p: p.name)
    if not xlsx_paths:
        raise FileNotFoundError(f"No XLSX files matched {args.pattern} under {out_dir}")

    for xlsx_path in xlsx_paths:
        codes = _extract_codes_from_xlsx(xlsx_path)
        csv_path = xlsx_path.with_suffix(".csv")
        _write_hts_csv(csv_path, codes, encoding=args.encoding)
        print(f"[csv] wrote: {csv_path} ({len(codes)} codes)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
