# -*- coding: utf-8 -*-
"""Build trading_statistics_events.xlsx from the two event CSVs.

- Sheet per case (case1_earnings / case2_news_gap), rows newest-first (CSV order).
- Rows that appear in ANY top-20% file -> light sky-blue row fill
  (+ helper column `top20_hit` listing which metrics).
- `turnover` / `turnover_fmt` cells -> 10-step green scale by decile
  (진할수록 거래대금 큼; 상위 3단계는 흰 글자).
- Return columns (gap/cc/o2c 1d/o2c 2d/fwd 1m/fwd 2m) -> 5-step diverging fill:
  >=+10% 짙은 초록, +3~10% 연초록, ±3% 중립, -10~-3% 연빨강, <=-10% 짙은 빨강.
"""
import glob
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, "..", "raw_data")
T20 = os.path.join(RAW, "top20pct_by_mcap")
OUT = os.path.join(RAW, "trading_statistics_events.xlsx")

RET_COLS = ["gap_pct", "cc_pct", "o2c_1d_pct", "o2c_2d_pct", "fwd_1m_pct", "fwd_2m_pct"]

# 10-step green ramp (light -> dark) for turnover deciles
GREENS = ["E9F7EC", "D4EFDA", "BFE7C8", "A9DFB6", "8FD3A0",
          "6FC287", "4FAF6D", "339955", "1F7E42", "115C2E"]
WHITE_FONT_FROM = 7            # deciles 8,9,10 get white text
# 5-step diverging for returns
RET_FILLS = {
    "dgreen": ("1F7E42", True),   # >= +10
    "lgreen": ("C9E9D2", False),  # +3 ~ +10
    "neutral": (None, False),     # -3 ~ +3  (no fill)
    "lred": ("F5C8C4", False),    # -10 ~ -3
    "dred": ("B93A32", True),     # <= -10
}
BLUE_ROW = "DDEBF7"
HEADER_FILL = "EDEDE9"


def ret_bucket(v):
    if pd.isna(v):
        return None
    if v >= 10:
        return "dgreen"
    if v >= 3:
        return "lgreen"
    if v > -3:
        return "neutral"
    if v > -10:
        return "lred"
    return "dred"


def load_top20_hits(case_key):
    """(ticker, event_date) -> 'metric|metric' from the per-bucket top20 files."""
    hits = {}
    for f in glob.glob(os.path.join(T20, case_key, "*.csv")):
        metric = os.path.basename(f).split("_top20_")[0]
        t = pd.read_csv(f, usecols=["ticker", "event_date"])
        for k in zip(t.ticker, t.event_date):
            hits.setdefault(k, set()).add(metric)
    return {k: "|".join(sorted(v)) for k, v in hits.items()}


wb = Workbook()
wb.remove(wb.active)

for case_key, fn, sheet in [("case1", "case1_earnings_events.csv", "case1_earnings"),
                            ("case2", "case2_news_gap_events.csv", "case2_news_gap")]:
    df = pd.read_csv(os.path.join(RAW, fn))
    hits = load_top20_hits(case_key)
    df["top20_hit"] = [hits.get((t, d), "") for t, d in zip(df.ticker, df.event_date)]

    decile = pd.qcut(df.turnover.rank(method="first"), 10, labels=False)

    ws = wb.create_sheet(sheet)
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center")

    col_idx = {c: i + 1 for i, c in enumerate(df.columns)}
    tno_cols = [col_idx[c] for c in ("turnover", "turnover_fmt") if c in col_idx]
    ret_cols = [(c, col_idx[c]) for c in RET_COLS if c in col_idx]

    for r, row in enumerate(df.itertuples(index=False), start=2):
        vals = list(row)
        is_top = bool(vals[col_idx["top20_hit"] - 1])
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=None if pd.isna(v) else v)
            if is_top:
                cell.fill = PatternFill("solid", fgColor=BLUE_ROW)
        # turnover decile greens (overrides row blue on those cells)
        d = decile.iloc[r - 2]
        if pd.notna(d):
            g = GREENS[int(d)]
            for c in tno_cols:
                cell = ws.cell(row=r, column=c)
                cell.fill = PatternFill("solid", fgColor=g)
                if int(d) >= WHITE_FONT_FROM:
                    cell.font = Font(color="FFFFFF")
        # diverging return fills
        for name, c in ret_cols:
            b = ret_bucket(vals[c - 1])
            if b and b != "neutral":
                hexv, white = RET_FILLS[b]
                cell = ws.cell(row=r, column=c)
                cell.fill = PatternFill("solid", fgColor=hexv)
                if white:
                    cell.font = Font(color="FFFFFF")

    # cosmetics: freeze header, autofilter, widths, number formats
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c, name in enumerate(df.columns, start=1):
        L = get_column_letter(c)
        if name.startswith("news_titles"):
            ws.column_dimensions[L].width = 60
        elif name in ("ticker", "event_date", "date_2d", "earnings_date"):
            ws.column_dimensions[L].width = 11
        else:
            ws.column_dimensions[L].width = max(9, min(len(name) + 2, 22))
        fmt = None
        if name in RET_COLS or name in ("sd60_gap", "sd60_cc", "z_gap", "z_cc", "gap_pctile"):
            fmt = "0.00"
        elif name in ("turnover", "pre_turnover_20d_median", "market_cap", "market_cap_at_event"):
            fmt = "#,##0"
        elif name.startswith(("prev_close", "open_", "close_")):
            fmt = "0.00"
        if fmt:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=c).number_format = fmt

    n_top = (df.top20_hit != "").sum()
    print(f"{sheet}: {len(df)}행, top20 표시 {n_top}행")

wb.save(OUT)
print("wrote", OUT)
